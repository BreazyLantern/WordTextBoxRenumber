from pathlib import Path
import os
from datetime import datetime
import sys
import lxml.etree as etree
from docx import Document
from docx.shared import Pt
import re

class FileManip:
    def __init__(self):
        self.working_file = None
        self.list_documents = []
        try:
            if getattr(sys, 'frozen', False):
                # This block is executed if the script is running inside a PyInstaller/similar executable
                print("Running in a standalone executable (.exe)")
                print(f"Executable path: {sys.executable}")
                cd = Path(__file__).parent if "__file__" in locals() else Path.cwd()
                file_path = os.path.abspath(cd)
                self.working_file = file_path
            else:
                # This block is executed if the script is running as a normal .py file
                print("Running as a Python script (.py)")
                print(f"Interpreter path: {sys.executable}")
                self.working_file = os.path.dirname(os.path.realpath(__file__))
        except Exception as e:
            input(f"Something went wrong with initialization: {e}")

    def output_paths(self):
        """Debugging variable inputs"""
        print(f"Current working directory {self.working_file}")
        if len(self.list_documents) != 0:
            print("Working files list in full path in local directory:")
            for doc in self.list_documents:
                print(doc)
        print()

    def add_to_list(self, new_list):
        self.list_documents.clear()
        for item in new_list:
            if ".doc" in Path(item).suffix:
                self.list_documents.append(item)

    def edit_text_in_documents(self, filename):
        doc = Document(filename)
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Ariel'
        font.size = Pt(20)
        for text in doc.textboxes:
            # add 6 to the texts
            new_text = f"{int(text.paragraphs[0].text) + 6}"
            text.paragraphs[0].text = new_text

        self.save_document(doc, filename)

    @staticmethod
    def get_unique_filename(base_filename, date, extension=".docx"):
        """
        Generates a unique filename by appending (1), (2), etc.
        if the base filename already exists.
        """
        filename = f"{base_filename}{date}{extension}"
        counter = 1
        while Path(filename).exists():
            filename = f"{base_filename}{date}({counter}){extension}"
            counter += 1
        return filename

    def save_document(self, doc, filename):
        date = f"_{datetime.today().strftime("%d %b, %Y")}"
        file_path = Path(filename).absolute()
        print(f"Working on file: {file_path.stem}")
        temp = None

        if file_path.stem.find("_") != -1:
            index = file_path.stem.find("_")

            temp = self.get_unique_filename(file_path.stem[:index], date, file_path.suffix)
        else:
            temp = self.get_unique_filename(file_path.stem, date, file_path.suffix)
        try:
            doc.save(temp )
            print(f"Document edited and saved to: {temp}")
            print()
        except Exception as e:
            print(f"Save system had an error: {e}")


    def work_on_all_files(self):
        if len(self.list_documents) != 0:
            for doc in self.list_documents:
                try:
                    self.edit_text_in_documents(doc)
                except Exception as e:
                    print(e)
        else:
            user_input = "Placeholder"
            while user_input != "":
                user_input = input("What Documents would you like to edit: ")
                print()
                docx_files = re.split(r' (?=(?:[^"]*"[^"]*")*[^"]*$)', user_input)
                new_list = []
                for docx in docx_files:
                    docx = docx.lstrip().replace('"', '').rstrip()
                    if Path(docx).is_file():
                        new_list.append(docx)

                self.add_to_list(new_list)
                for doc in self.list_documents:
                    try:
                        self.edit_text_in_documents(doc)
                    except Exception as e:
                        print(e)
